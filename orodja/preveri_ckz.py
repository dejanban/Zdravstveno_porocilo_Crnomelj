"""Preveri gradivo CKZ in izdela njegov PDF; uporablja le prazne predloge."""
from pathlib import Path
import json
import fitz
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

P = Path(__file__).resolve().parent.parent
OUT = P / 'podatki/ckz'


def main():
    (P / 'preverjanje').mkdir(exist_ok=True)
    errors, requests = [], []
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True)
        page = browser.new_page()
        page.on('pageerror', lambda e: errors.append(str(e)))
        page.on('request', lambda r: requests.append(r.url))
        page.goto((OUT / 'ckz-obrazci.html').as_uri())
        assert page.locator('input,textarea,form').count() == 0
        assert page.locator('h1').count() == 4
        for width in [320, 390, 768, 1440, 1920]:
            page.set_viewport_size({'width': width, 'height': 1000})
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth+1'), width
        page.emulate_media(media='print')
        page.pdf(path=str(OUT / 'ckz-obrazci.pdf'), format='A4', prefer_css_page_size=True, print_background=True)
        page.emulate_media(media='screen')
        page.goto((P / 'porocilo.html').as_uri())
        assert page.locator('#ckz-spremljanje').count() == 1
        for width in [320, 390, 768, 1440, 1920]:
            page.set_viewport_size({'width': width, 'height': 1000})
            page.locator('#ckz-spremljanje').scroll_into_view_if_needed()
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth+1'), width
            if width in (390, 1440):
                page.screenshot(path=str(P / f'preverjanje/ckz-spremljanje-{width}.png'))
        browser.close()
    assert not errors, errors
    assert not any(u.startswith(('http:', 'https:')) for u in requests), requests
    soup = BeautifulSoup((OUT / 'ckz-obrazci.html').read_text(encoding='utf-8'), 'html.parser')
    for link in soup.select('a[href]'):
        href = link['href']
        if href.startswith(('http:', 'https:')): continue
        path = (OUT / href.split('#')[0]).resolve()
        assert path.is_relative_to(P) and path.exists(), href
    wb = load_workbook(OUT / 'ckz-spremljanje.xlsx')
    assert len(wb.sheetnames) == 5
    for row in wb['Kazalniki'].iter_rows(min_row=2):
        assert row[3].value is None and row[4].value is None
        assert row[5].data_type == 'f'
    assert all(wb['Samoocena ekipe'].cell(i, 2).value is None for i in range(2, 10))
    doc = fitz.open(OUT / 'ckz-obrazci.pdf')
    assert len(doc) == 4, f'Obrazci morajo imeti štiri strani, ne {len(doc)}.'
    assert 'vprašalnik za udeleženca' in doc[1].get_text()
    assert 'Literatura' in doc[3].get_text()
    for i in range(len(doc)):
        doc[i].get_pixmap(matrix=fitz.Matrix(1.2, 1.2)).save(P / f'preverjanje/ckz-obrazci-stran-{i+1}.png')
    result = {'strani_obrazcev_pdf': len(doc), 'listi_xlsx': len(wb.sheetnames), 'preverjanje': 'Prazne predloge, lokalne povezave, 5 širin, brez omrežja in napak JavaScript.'}
    (P / 'preverjanje/ckz.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
