"""Celovit lokalni tehnični pregled; virov in pričakovanih vsot ne spreminja."""
from pathlib import Path
from collections import Counter
from datetime import datetime, timezone
import csv
import hashlib
import json
import math
import re
import sys
import xml.etree.ElementTree as ET
from urllib.parse import unquote, urlsplit
import fitz
from bs4 import BeautifulSoup
from openpyxl import load_workbook

P = Path(__file__).resolve().parent.parent


def main():
    errors, warnings, counts = [], [], Counter()
    def issue(kind, path, detail):
        kind.append({'pot': str(path), 'opis': str(detail)})
    def read(path):
        return json.loads((P / path).read_text(encoding='utf-8-sig'))
    files = [p for root in ('podatki', 'viri', 'wiki', 'vsebina')
             for p in (P / root).rglob('*') if p.is_file() and not p.name.startswith('~$')]
    parsed = {}
    for p in files:
        rel = p.relative_to(P).as_posix()
        counts['datoteke'] += 1
        try:
            if p.stat().st_size == 0:
                issue(warnings, rel, 'Prazna datoteka.')
            if p.suffix in ('.json', '.geojson'):
                def pairs(items):
                    d = {}
                    for k, v in items:
                        if k in d: raise ValueError('Podvojen ključ JSON: ' + k)
                        d[k] = v
                    return d
                def constant(value):
                    raise ValueError('Nestandardna številka JSON: ' + value)
                parsed[rel] = json.loads(p.read_text(encoding='utf-8-sig'), object_pairs_hook=pairs, parse_constant=constant)
                counts['json'] += 1
            elif p.suffix == '.csv':
                with p.open(encoding='utf-8-sig', newline='') as f:
                    sample = f.read(8192); f.seek(0)
                    dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
                    rows = list(csv.reader(f, dialect))
                for n, row in enumerate(rows[1:], 2):
                    if row and len(row) != len(rows[0]):
                        issue(errors, rel, f'Vrstica {n}: {len(row)} polj namesto {len(rows[0])}.')
                counts['csv'] += 1
            elif p.suffix in ('.gpx', '.svg', '.xml'):
                tree = ET.parse(p)
                if p.suffix == '.gpx':
                    for el in tree.iter():
                        if el.tag.split('}')[-1] in ('trkpt', 'rtept', 'wpt'):
                            lat, lon = float(el.attrib['lat']), float(el.attrib['lon'])
                            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                                raise ValueError('Neveljavne koordinate GPX.')
                counts[p.suffix[1:]] += 1
            elif p.suffix == '.xlsx':
                wb = load_workbook(p, read_only=True, data_only=True)
                for sheet in wb:
                    for row in sheet:
                        for cell in row:
                            if cell.data_type == 'e':
                                issue(warnings, rel, f'Excelova napaka {sheet.title}!{cell.coordinate}: {cell.value}')
                wb.close(); counts['xlsx'] += 1
            elif p.suffix == '.pdf':
                with fitz.open(p) as doc:
                    if not len(doc): raise ValueError('PDF brez strani.')
                    for page in doc: page.get_text()
                counts['pdf'] += 1
        except Exception as exc:
            issue(errors, rel, exc)
    manifest = read('viri/manifest.json')['datoteke']
    registered = set()
    for row in manifest:
        name = row['lokalna_pot']; p = (P / name).resolve()
        if name in registered: issue(errors, name, 'Podvojen vnos manifesta.')
        registered.add(name)
        if not p.is_relative_to(P) or not p.is_file():
            issue(errors, name, 'Vir manjka ali zapušča projekt.'); continue
        actual = hashlib.sha256(p.read_bytes()).hexdigest()
        if actual != row['sha256']:
            issue(errors, name, f"SHA-256: pričakovano {row['sha256']}; dejansko {actual}")
        else: counts['sha256_ujemanja'] += 1
    counts['manifest'] = len(manifest)
    for p in (P / 'viri').rglob('*'):
        if p.is_file() and not p.name.startswith('~$') and p.suffix.lower() in ('.pdf', '.xlsx', '.xls', '.gpx') and p.relative_to(P).as_posix() not in registered:
            issue(warnings, p.relative_to(P), 'Izvirnik ni v manifestu SHA-256.')
    for ledger in [P/'podatki/literatura/literatura.json', *(P/'podatki').rglob('*-viri.json')]:
        data = read(ledger.relative_to(P))
        rows = data if isinstance(data, list) else data.values()
        for row in rows:
            if isinstance(row, dict):
                for key in ('local', 'lokalna_kopija', 'lokalna_pot'):
                    if row.get(key):
                        target = (P/row[key]).resolve()
                        if not target.is_relative_to(P) or not target.exists():
                            issue(errors, ledger.relative_to(P), 'Manjkajoča lokalna kopija: '+row[key])
                        counts['bibliografske_poti'] += 1
    records = read('podatki/register/register-dogodkov.json')
    counts['objave_registra'] = len(records)
    if len({r['url'] for r in records}) != len(records): issue(errors, 'register', 'Podvojeni URL-ji.')
    for row in records:
        if not (P/row['lokalna_kopija']).is_file(): issue(errors, 'register', row['lokalna_kopija'])
    stats = read('podatki/register/statistika-registra.json')
    if dict(Counter(r['pripadnost'] for r in records)) != stats['pripadnost']:
        issue(errors, 'register', 'Statistika pripadnosti se ne ujema z registrom.')
    with (P/'podatki/register/register-dogodkov.csv').open(encoding='utf-8-sig', newline='') as f:
        exported = list(csv.DictReader(f, delimiter=';'))
    if exported != [{k: str(v) if v is not None else '' for k,v in r.items()} for r in records]:
        issue(errors, 'register', 'CSV in JSON se ne ujemata po vseh poljih.')
    for key, item in read('podatki/nijz/kazalniki.json').items():
        years = item['years']
        if years != sorted(set(years)): issue(errors, key, 'Leta niso enolična in urejena.')
        for series in item['series']:
            if len(years) != len(series['values']): issue(errors, key, 'Dolžina niza ne ustreza letom.')
            if any(v is not None and (not isinstance(v, (int,float)) or not math.isfinite(v)) for v in series['values']):
                issue(errors, key, 'Nenumerična vrednost kazalnika.')
        counts['kazalniki_nijz'] += 1
    routes = read('podatki/poti/poti-podatki.json')['poti']
    counts['poti'] = len(routes)
    if len(routes) != len(list((P/'viri/gpx').rglob('*.gpx'))): issue(errors, 'poti', 'Število GPX se ne ujema.')
    for row in routes:
        if hashlib.sha256((P/row['datoteka']).read_bytes()).hexdigest() != row['sha256']:
            issue(errors, row['datoteka'], 'SHA-256 se ne ujema s podatki poti.')
        if not 0 <= row['km_v_obcini'] <= row['km'] + 1e-6 or row['neveljavne_tocke']:
            issue(errors, row['datoteka'], 'Neveljavna dolžina ali točke poti.')
    aed = read('podatki/aed/aed.json')['aed']
    wb = load_workbook(P/'viri/AEDs.xlsx', read_only=True, data_only=True)
    source = [r for r in list(wb.active.values)[1:] if r[1]]
    counts['aed_izvirnik'] = len(source)
    counts['aed_json'] = len(aed)
    if len(source) != len(aed): issue(errors, 'AED', 'Število zapisov XLSX in JSON se ne ujema.')
    elif any((float(r[1]),float(r[2]),str(r[3] or '').removeprefix('AED ').strip(),str(r[4] or '').strip()) != (a['lon'],a['lat'],a['ime'],a['naslov']) for r,a in zip(source,aed)):
        issue(errors, 'viri/AEDs.xlsx', 'Koordinate, imena ali naslovi izvirnika se ne ujemajo z aed.json.')
    wb.close()
    with (P/'podatki/aed/aed.csv').open(encoding='utf-8-sig', newline='') as f:
        reader = csv.DictReader(f, delimiter=';'); rows = list(reader)
        expected = ['ime','naslov','lon','lat','v_obcini_crnomelj']
        if reader.fieldnames != expected:
            issue(errors, 'podatki/aed/aed.csv', 'Glava CSV ne ustreza generatorju: '+str(reader.fieldnames))
        bad = []
        for n,row in enumerate(rows,2):
            try:
                lon,lat = float(row['lon']),float(row['lat'])
                if not (-180<=lon<=180 and -90<=lat<=90): raise ValueError()
            except (ValueError,KeyError): bad.append(n)
        if bad: issue(errors,'podatki/aed/aed.csv','Neveljavne koordinate v vrsticah: '+str(bad))
    documents = list((P/'viri').rglob('*.html')) + list((P/'podatki').rglob('*.html')) + list(P.glob('*.html'))
    for p in documents:
        soup = BeautifulSoup(p.read_text(encoding='utf-8'), 'html.parser')
        ids = {e['id'] for e in soup.select('[id]')}
        counts['html'] += 1
        for el in soup.select('[href],[src]'):
            u = el.get('href') or el.get('src')
            parts = urlsplit(u)
            if parts.scheme in ('http','https','mailto','tel','data') or parts.netloc: continue
            target = (p.parent/unquote(parts.path)).resolve() if parts.path else p
            if parts.scheme or not target.is_relative_to(P) or not target.exists():
                issue(errors, p.relative_to(P), 'Neveljavna lokalna povezava: '+u)
            elif not parts.path and parts.fragment and unquote(parts.fragment) not in ids:
                issue(errors, p.relative_to(P), 'Manjkajoče sidro: '+u)
            counts['lokalne_povezave'] += 1
    for p in (P/'wiki').rglob('*.md'):
        for link in re.findall(r'\[\[([^\]]+)\]\]', p.read_text(encoding='utf-8')):
            name = link.split('|')[0].rstrip('\\').split('#')[0]
            if not name: continue
            target = (p.parent/name).resolve()
            if not target.exists() and not target.suffix: target = target.with_suffix('.md')
            if not target.is_relative_to(P) or not target.exists():
                issue(warnings, p.relative_to(P), 'Nerazrešena wikipovezava: '+name)
            counts['wiki_povezave'] += 1
    result = {'datum':datetime.now(timezone.utc).isoformat(), 'obseg':'Lokalna tehnična preverba, brez nove vsebinske ali spletne verifikacije; brez preračuna formul in vizualne preverbe.', 'stevci':dict(counts), 'napake':errors, 'opozorila':warnings}
    (P/'preverjanje/podatki-pregled.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'stevci':dict(counts),'napake':len(errors),'opozorila':len(warnings)},ensure_ascii=False))
    return bool(errors)


if __name__ == '__main__':
    sys.exit(main())
