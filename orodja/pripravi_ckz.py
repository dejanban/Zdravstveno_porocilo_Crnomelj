"""Prazni obrazci CKZ in preglednica agregatov; brez osebnih podatkov."""
from pathlib import Path
import markdown
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

P = Path(__file__).resolve().parent.parent
OUT = P / 'podatki/ckz'


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    source = (P / 'vsebina/ckz-obrazci.md').read_text(encoding='utf-8')
    source = re.sub(r'_{3,}', lambda m: '<span class="blank ' + ('short' if len(m[0]) < 10 else 'long') + '">&nbsp;</span>', source)
    parts = source.split('<div class="prelom"></div>')
    body = ''.join(f'<section class="sheet sheet-{i}">' + markdown.markdown(part, extensions=['tables']) + '</section>' for i, part in enumerate(parts, 1))
    css = '''*{box-sizing:border-box}body{font:16px/1.55 Arial,sans-serif;color:#163632;background:#f4f5ef;margin:0}main{max-width:1050px;margin:auto;padding:30px}h1{font-size:30px;line-height:1.15}h2{font-size:22px}table{border-collapse:collapse;width:100%;margin:16px 0;font-size:14px}th,td{border:1px solid #cdd8cc;padding:10px;text-align:left;vertical-align:top;overflow-wrap:anywhere}th{background:#e6eddf}a{color:#17695b}button{font:inherit;padding:8px 14px;cursor:pointer}.toolbar{display:flex;gap:20px;flex-wrap:wrap;margin-bottom:24px}.prelom{border-top:2px solid #17695b;margin:35px 0}a:focus-visible,button:focus-visible{outline:3px solid #a65c20}@media(max-width:600px){main{padding:16px}h1{font-size:26px}table{font-size:12px}td,th{padding:7px}}@media print{@page{size:A4;margin:12mm}body{background:white;font-size:9pt;line-height:1.3}main{padding:0;max-width:none}.toolbar{display:none}h1{font-size:19pt;margin:0 0 10px}h2{font-size:13pt}table{font-size:8.5pt;margin:9px 0}th,td{padding:6px}p{margin:9px 0}tr{break-inside:avoid}thead{display:table-header-group}.prelom{break-before:page;border:0;margin:0}a{overflow-wrap:anywhere}*{print-color-adjust:exact;-webkit-print-color-adjust:exact}}'''
    css += '.blank{display:inline-block;border-bottom:1px solid #75877d;min-height:1em}.blank.short{width:3em;max-width:100%}.blank.long{width:100%;min-width:4em}.sheet+.sheet{border-top:2px solid #17695b;margin-top:35px;padding-top:25px}.sheet-1 table:first-of-type,.sheet-1 table:last-of-type{table-layout:fixed}.sheet-1 table:nth-of-type(2) th:nth-child(1){width:46%}.sheet-1 table:nth-of-type(2) th:nth-child(2){width:16%}.sheet-1 table:nth-of-type(2) th:nth-child(3){width:38%}.sheet-4 table:nth-of-type(2) th:first-child{width:20%}@media print{.sheet+.sheet{break-before:page;border:0;margin:0;padding:0}.sheet-1 table,.sheet-2 table{font-size:9pt}.sheet-1 td,.sheet-2 td{padding:7px}.sheet-1 table:last-of-type td{height:34px}}'
    page = f'''<!doctype html><html lang="sl"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>CKZ · obrazci za spremljanje napredka</title><style>{css}</style></head><body><main><nav class="toolbar" aria-label="Gradivo"><a href="../../porocilo.html#ckz-spremljanje">← Poročilo</a><a href="ckz-obrazci.pdf">Prenesi PDF</a><a href="ckz-spremljanje.xlsx">Excelova predloga</a><button onclick="window.print()">Natisni obrazce</button></nav>{body}</main></body></html>'''
    (OUT / 'ckz-obrazci.html').write_text(page, encoding='utf-8')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Navodila'
    for row in [
        ['CKZ · predloga agregatov 1.0', '22. 9. 2026 · prazna delovna predloga'],
        ['Namen', 'Vpisujte le zbirne podatke. Brez imen, šifer oseb, kontaktov ali individualnih odgovorov.'],
        ['Obdobje', 'Vsaka vrstica Kazalnikov je svoj presek istega programa in jasno opredeljene kohorte. Vrstice za nov presek kopirajte.'],
        ['Vnos', 'Modra polja so vnosna. Prazno = ni podatka; 0 = preverjena ničla. Rezultati se preračunajo ob odprtju v Excelu/LibreOffice.'],
        ['Kontrole', 'Za T6/T12 vključite vse začetnike, tudi nedokončane, katerih okno ±30 dni je do preseka že poteklo.'],
        ['Ne seštevajte', 'Različnih oseb po mesecih/programih ne seštevajte v letno število. Kohort z različnimi datumi ne mešajte.'],
        ['Izidi', 'Parne spremembe izračuna CKZ v varovanem okolju; sem vpiše samo agregate za isto navado.'],
        ['Samoocena', 'Lestvica 1–5 v obrazcih; N/P ni podatka, N/U ni uporabno. Delno povprečje ni primerljivo s celotnim.'],
        ['Objava', 'Pred objavo preverite majhne celice in izpeljavo iz seštevkov; ne objavite tega delovnega zvezka avtomatično.'],
        ['Cilji', 'Dogovorite jih po začetnem pregledu. Predloga nima vnaprej vpisanih rezultatov ali pragov uspešnosti.'],
        ['Izvor', 'Avtorski predlog; navodila in vir NIJZ so v ckz-obrazci.html/PDF. Ni uradni obrazec NIJZ.'],
    ]: ws.append(row)
    k = wb.create_sheet('Kazalniki')
    k.append(['Kazalnik', 'Program / kohorta', 'Presek (datum)', 'Števec', 'Imenovalec', 'Rezultat (%)', 'Pravilo', 'Cilj (%)', 'Rok', 'Odgovorna vloga', 'Opomba / ukrep'])
    definitions = [
        ('Začetek v 30 dneh', 'Začeli v 30 dneh / napoteni iste kohorte z vsaj 30 dnevi opazovanja.'),
        ('Zaključek programa', 'Zaključili / vsi začetniki iste kohorte z zapadlim predvidenim zaključkom.'),
        ('Odziv T6', 'Odgovorili / vsi začetniki z do preseka zaprtim oknom T6; tudi nedokončani.'),
        ('Odziv T12', 'Odgovorili / vsi začetniki z do preseka zaprtim oknom T12; tudi nedokončani.'),
        ('Pari navade T0–T6', 'Veljavni pari za isto navado / vsi začetniki z zaprtim oknom T6.'),
        ('Pari navade T0–T12', 'Veljavni pari za isto navado / vsi začetniki z zaprtim oknom T12.'),
        ('Zadovoljstvo 4 ali 5', 'Odgovori 4 ali 5 / veljavni odgovori 1–5 na vprašanje 9 ob zaključku.'),
        ('Odziv na zadovoljstvo', 'Veljavni odgovori 1–5 / vsi povabljeni k vprašanju 9 v isti kohorti.'),
        ('Izvedeni ukrepi', 'Izvedeni ukrepi do roka / vsi ukrepi z zapadlim rokom v obdobju.'),
    ]
    for i, (name, rule) in enumerate(definitions, 2):
        formula = f'=IF(OR(D{i}="",E{i}=""),"ni podatka",IF(OR(D{i}<0,E{i}<0,D{i}>E{i}),"preveri vnos",IF(E{i}=0,"ni izračunljivo",D{i}/E{i})))'
        k.append([name, None, None, None, None, formula, rule])
        k.cell(i, 6).number_format = '0.0%'
        k.cell(i, 8).number_format = '0.0%'
    dv = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=0, allow_blank=True)
    dv.errorTitle = 'Preverite število'; dv.error = 'Vpišite nenegativno celo število ali pustite prazno.'; dv.showErrorMessage = True
    k.add_data_validation(dv); dv.add('D2:E10')
    c = wb.create_sheet('Obseg in izidi')
    c.append(['Podatek', 'Program / kohorta / ista navada', 'Obdobje / presek', 'Vrednost', 'Enota', 'Metoda / vir / manjkajoče', 'Cilj', 'Rok / odgovorna vloga'])
    for name, unit in [
        ('Različne osebe v obdobju', 'osebe'), ('Obiski v obdobju', 'obiski'),
        ('Kontrole še niso zapadle T6', 'osebe'), ('Kontrole še niso zapadle T12', 'osebe'),
        ('Brez odgovora T6', 'osebe'), ('Brez odgovora T12', 'osebe'),
        ('Mediana parnih razlik navade T6 − T0', 'dni/teden'), ('Veljavni pari za to mediano T6', 'pari'),
        ('Mediana parnih razlik navade T12 − T0', 'dni/teden'), ('Veljavni pari za to mediano T12', 'pari'),
        ('Mediana čakanja med napotitvijo in začetkom', 'dni; samo začeli'),
        ('Število oseb, ki še čakajo ob preseku', 'osebe'),
        ('Pripisani stroški kohorte', 'EUR'), ('Vsi začetniki iste kohorte', 'osebe'), ('Vsi zaključili isto kohorto', 'osebe'),
        ('Porabljene ure ekipe za spremljanje', 'ure'),
    ]: c.append([name, None, None, None, unit])
    s = wb.create_sheet('Samoocena ekipe')
    s.append(['Merilo', 'Ocena 1–5 / N/P / N/U', 'Dokazilo / datum / razlog', 'Ukrep', 'Rok', 'Odgovorna vloga'])
    for label in ['Preglednost', 'Dostop', 'Ujemanje s potrebami', 'Dokazi in izidi', 'Pravičnost', 'Stroški', 'Trajnost', 'Izvedljivost']: s.append([label])
    s.append(['Število ocenjenih meril (od 8)', '=COUNT(B2:B9)'])
    s.append(['Povprečje ocenjenih meril', '=IF(COUNT(B2:B9)=0,"ni podatka",AVERAGE(B2:B9))'])
    s.append(['Popolnost', '=IF(COUNT(B2:B9)=8,"celotno povprečje","delno; ne primerjajte različnih naborov")'])
    s.append(['Program / obdobje / datum pregleda'])
    dv = DataValidation(type='list', formula1='"1,2,3,4,5,N/P,N/U"', allow_blank=True)
    dv.showErrorMessage = True; s.add_data_validation(dv); dv.add('B2:B9')
    u = wb.create_sheet('Dnevnik ukrepov')
    u.append(['Datum pregleda', 'Program', 'Težava in dokaz', 'Ukrep', 'Odgovorna vloga', 'Rok', 'Izhodišče in cilj', 'Datum preverbe', 'Rezultat / nadaljnji sklep'])
    for _ in range(12): u.append([None] * 9)
    for sheet in wb:
        sheet.freeze_panes = 'B2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='17695B')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.value is None: cell.fill = PatternFill('solid', fgColor='E7F1FA')
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = 36 if letter == 'A' else 28
        sheet.row_dimensions[1].height = 36
        for i in range(2, sheet.max_row + 1): sheet.row_dimensions[i].height = 62
    ws.column_dimensions['B'].width = 110
    k.column_dimensions['G'].width = 75
    wb.save(OUT / 'ckz-spremljanje.xlsx')
    print('CKZ: pripravljena prazna obrazca HTML in preglednica XLSX.')


if __name__ == '__main__':
    main()
